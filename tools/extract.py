#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从《检验项目手册》V3.0.xlsx 抽取数据，生成 assets/data.js（浏览器可直接加载，file:// 亦可）。

表结构：A 列为医嘱名称（合并单元格纵向跨行），每个医嘱占一段行；
段内每行是一条子项（B 子项名称 / C 对应收费 / D 价格），医嘱级属性在段首行。

用法：python3 tools/extract.py ["../《检验项目手册》V3.0.xlsx"]
"""
import collections
import difflib
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

DEFAULT_XLSX = os.path.join(os.path.dirname(ROOT), "《检验项目手册》V3.0.xlsx")

COLS = ["order_name", "sub_name", "charge", "price", "total", "abbr",
        "category", "kind", "specimen", "tube", "turnaround",
        "lab", "phone", "note", "active"]


def clean(v):
    """去空白、去不可见字符；空值统一为 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).replace("\u3000", " ").replace("\xa0", " ").strip()
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\s*\n\s*", "\n", s)
    return s or None


def to_price(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace(",", "").replace("，", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    return round(float(m.group()), 2)


def build_grid(ws):
    """展开合并单元格，返回二维数组 grid[row][col]，row/col 从 1 开始（索引 0 空占位）。"""
    merged = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = v
    grid = {}
    for r in range(2, ws.max_row + 1):
        grid[r] = [None] + [merged[(r, c)] if (r, c) in merged else ws.cell(r, c).value
                            for c in range(1, len(COLS) + 1)]
    return grid


def a_column_starts(ws):
    """A 列中"医嘱名称"的起始行（合并单元格只取首行）。"""
    continued = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1 and rng.max_col == 1 and rng.max_row > rng.min_row:
            for r in range(rng.min_row + 1, rng.max_row + 1):
                continued.add(r)
    starts = []
    for r in range(2, ws.max_row + 1):
        if clean(ws.cell(r, 1).value) and r not in continued:
            starts.append(r)
    return starts


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    grid = build_grid(ws)
    starts = a_column_starts(ws)

    entries = []
    seen_ids = collections.Counter()
    warnings = []

    for i, start in enumerate(starts):
        end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_row
        head = grid[start]
        meta = {name: clean(head[idx + 1]) for idx, name in enumerate(COLS)}

        items = []
        for r in range(start, end + 1):
            row = grid[r]
            sub = clean(row[2])          # B 子项名称
            charge = clean(row[3])       # C 对应收费
            price = to_price(row[4])     # D 价格
            if not (sub or charge):
                continue
            items.append({"name": sub or charge, "charge": charge, "price": price})

        total = to_price(head[5])        # E 总价
        sum_price = round(sum(it["price"] or 0 for it in items), 2)

        kind = meta["kind"] or ("组合" if len(items) > 1 else "单项")

        entry = {
            "id": None,
            "name": meta["order_name"],
            "kind": kind,
            "category": meta["category"],
            "specimen": meta["specimen"],
            "tube": meta["tube"],
            "turnaround": meta["turnaround"],
            "lab": meta["lab"],
            "phone": meta["phone"],
            "note": meta["note"],
            "active": meta["active"],
            "total": total if total is not None else sum_price,
            "itemCount": len(items),
            "items": items,
            "row": start,
        }

        base = "p%d" % start
        seen_ids[base] += 1
        entry["id"] = base if seen_ids[base] == 1 else "%s-%d" % (base, seen_ids[base])
        entries.append(entry)

        # 数据体检（只打印，不写入产物）
        if total is not None and items and abs(total - sum_price) > 0.01:
            warnings.append("行%d【%s】总价%s ≠ 子项合计%s" % (start, entry["name"], total, sum_price))
        if kind == "组合" and len(items) < 2:
            warnings.append("行%d【%s】标记为组合但子项数=%d" % (start, entry["name"], len(items)))
        if kind == "单项" and len(items) != 1:
            warnings.append("行%d【%s】标记为单项但子项数=%d" % (start, entry["name"], len(items)))
        if not meta["category"]:
            warnings.append("行%d【%s】缺少医嘱分类" % (start, entry["name"]))
        if not meta["lab"]:
            warnings.append("行%d【%s】缺少责任组室" % (start, entry["name"]))

    # ---- 交叉索引：某个单项出现在哪些组合里 / 组合里的子项能链到哪个单项 ----
    # 注意：多个单项可能共用同一个"对应收费"名（血清/尿/脑脊液肌酐都收"肌酐(Cr)测定"），
    # 因此按 子项名称 -> 对应收费 的顺序匹配；候选不唯一时用标本类型消歧，仍不唯一则放弃链接（宁缺勿错）。
    singles = [e for e in entries if e["kind"] == "单项"]
    by_name = collections.defaultdict(list)     # 单项医嘱名 -> 单项
    by_charge = collections.defaultdict(list)   # 单项对应收费名 -> 单项
    for e in singles:
        by_name[normalize(e["name"])].append(e)
        for it in e["items"]:
            if it["charge"]:
                by_charge[normalize(it["charge"])].append(e)

    def pick(cands, combo):
        if not cands:
            return None
        # 先按标本类型过滤：尿液组合里的"肌酐"不应链到血清肌酐单项
        compat = [c for c in cands if specimen_kinds(combo["specimen"]) & specimen_kinds(c["specimen"])
                  or not specimen_kinds(combo["specimen"]) or not specimen_kinds(c["specimen"])]
        if not compat:
            return None
        if len(compat) == 1:
            return compat[0]
        # 标本描述不止一个候选时，取与组合标本描述最贴合的那个（如"24小时尿液"优先于"尿液"）
        scored = [(specimen_rank(combo["specimen"], c["specimen"]), c) for c in compat]
        best = max(s for s, _ in scored)
        top = [c for s, c in scored if s == best]
        return top[0] if best > (0, 0) and len(top) == 1 else None

    used_in = collections.defaultdict(set)     # 单项 id -> 组合 id 集合
    link = {}                                   # "组合id|子项序号" -> 单项 id
    ambiguous = 0
    for combo in entries:
        if combo["kind"] != "组合":
            continue
        for idx, it in enumerate(combo["items"]):
            target = pick(by_name.get(normalize(it["name"])), combo)
            if target is None:
                target = pick(by_charge.get(normalize(it["charge"])), combo)
            if target is not None:
                link["%s|%d" % (combo["id"], idx)] = target["id"]
                used_in[target["id"]].add(combo["id"])
            elif by_name.get(normalize(it["name"])) or by_charge.get(normalize(it["charge"])):
                ambiguous += 1

    for e in entries:
        e["usedIn"] = sorted(used_in.get(e["id"], []), key=lambda cid: next(
            (x["name"] for x in entries if x["id"] == cid), ""))
    for combo in entries:
        for idx, it in enumerate(combo["items"]):
            it["link"] = link.get("%s|%d" % (combo["id"], idx))

    payload = {
        "title": "检验项目手册",
        "version": "V3.0",
        "source": os.path.basename(xlsx),
        "generated": __import__("datetime").date.today().isoformat(),
        "stats": {
            "entries": len(entries),
            "combos": sum(1 for e in entries if e["kind"] == "组合"),
            "singles": sum(1 for e in entries if e["kind"] == "单项"),
        },
        "entries": entries,
    }

    out_dir = os.path.join(ROOT, "assets")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "data.js")
    with open(out, "w", encoding="utf-8") as f:
        f.write("/* 由 tools/extract.py 自动生成，请勿手工编辑。源文件：%s */\n" % payload["source"])
        f.write("window.HANDBOOK = ")
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")

    print("已生成 %s" % out)
    print("条目 %d（组合 %d / 单项 %d），子项总数 %d"
          % (payload["stats"]["entries"], payload["stats"]["combos"], payload["stats"]["singles"],
             sum(e["itemCount"] for e in entries)))
    print("子项可跳转链接：%d / %d（歧义未链接：%d）"
          % (len(link), sum(e["itemCount"] for e in entries if e["kind"] == "组合"), ambiguous))
    if warnings:
        print("\n数据体检提示 %d 条：" % len(warnings))
        for w in warnings:
            print("  -", w)


def normalize(s):
    """名称归一化：用于子项与单项之间的模糊匹配。"""
    if not s:
        return ""
    s = str(s).strip().lower()
    s = s.replace("（", "(").replace("）", ")").replace("，", ",")
    s = re.sub(r"[\s\u3000]+", "", s)
    s = re.sub(r"(测定|检测|试验|定量)$", "", s)
    s = re.sub(r"[()（）,，、\-—:：]", "", s)
    return s


SPECIMEN_GROUPS = {
    "blood": ("血清", "血浆", "全血"),
    "urine": ("尿",),
    "stool": ("粪便", "便盒"),
    "sputum": ("痰",),
    "swab": ("拭子",),
    "csf": ("脑脊液",),
    "fluid": ("积液", "胸水", "腹水", "穿刺液"),
    "tissue": ("组织",),
    "secretion": ("分泌物",),
    "balf": ("灌洗液",),
    "bone_marrow": ("骨髓",),
}


def specimen_kinds(specimen):
    """把标本类型描述映射成类别集合，用于判断组合与单项是否同源标本。"""
    s = specimen or ""
    return {kind for kind, keys in SPECIMEN_GROUPS.items() if any(k in s for k in keys)}


def specimen_rank(combo_spec, cand_spec):
    """标本贴合度排序键：描述完全一致优先，其次最长公共子串更长者优先。"""
    a, b = normalize(combo_spec), normalize(cand_spec)
    return (1 if a and a == b else 0, specimen_score(combo_spec, cand_spec))


def specimen_score(a, b):
    """两条标本描述的贴合度：最长公共子串长度（"24小时尿液"对"24小时尿液,硼酸防腐剂"得分更高）。"""
    a, b = normalize(a), normalize(b)
    if not a or not b:
        return 0
    return difflib.SequenceMatcher(None, a, b).find_longest_match(0, len(a), 0, len(b)).size


if __name__ == "__main__":
    main()
